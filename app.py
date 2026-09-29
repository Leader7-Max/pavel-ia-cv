"""Pavel IA CV — générateur de CV et lettres de motivation sur-mesure."""

import json
import re
import threading
from io import BytesIO
from collections import Counter
from pathlib import Path

import streamlit as st
from fpdf import FPDF

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    DOCX_AVAILABLE = True
except ImportError:  # python-docx absent : l'export Word est simplement désactivé
    DOCX_AVAILABLE = False
from google import genai
from google.genai import types

# ─────────────────────────────────────────────────────────────
# Configuration générale
# ─────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="Pavel IA CV",
    page_icon="📄",
    layout="centered",
    initial_sidebar_state="collapsed",
)

MODELS = ["gemini-2.5-flash", "gemini-2.0-flash", "gemini-3.5-flash-lite"]  # essayés dans l'ordre : les suivants servent de secours

TONES = [
    "Professionnel et formel (classique)",
    "Dynamique et moderne (impactant)",
    "Reconversion / profil atypique",
    "Audacieux et créatif",
]

LENGTHS = {
    "cv": {
        "Concise": "tient sur une page, 3 puces maximum par expérience",
        "Standard": "tient sur une page, 3 à 4 puces par expérience",
        "Détaillée": "peut occuper deux pages, 4 à 6 puces par expérience",
    },
    "lettre": {
        "Concise": "200 à 250 mots",
        "Standard": "300 à 350 mots",
        "Détaillée": "400 à 450 mots",
    },
}

LOGO_SVG = """<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 64 64" role="img" aria-label="Logo Pavel IA CV"><rect width="64" height="64" rx="16" fill="#0E6B6B"/><path d="M18 12h18l10 10v26a3 3 0 0 1-3 3H18a3 3 0 0 1-3-3V15a3 3 0 0 1 3-3z" fill="#FFFFFF"/><path d="M36 12v7a3 3 0 0 0 3 3h7z" fill="#B7D6D6"/><rect x="21" y="27" width="16" height="3" rx="1.5" fill="#0E6B6B"/><rect x="21" y="34" width="11" height="3" rx="1.5" fill="#9CC3C3"/><circle cx="44" cy="44" r="10" fill="#F2A93B" stroke="#0E6B6B" stroke-width="3"/><path d="M39.5 44.5l3.3 3.3 6-6.6" fill="none" stroke="#12272B" stroke-width="2.8" stroke-linecap="round" stroke-linejoin="round"/></svg>"""

CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Bricolage+Grotesque:opsz,wght@12..96,500;12..96,700&family=DM+Sans:wght@400;500;700&display=swap');

:root {
  --ink: #12272B;
  --teal: #0E6B6B;
  --amber: #F2A93B;
  --paper: #F5F7F8;
  --line: #DCE3E5;
  --muted: #5B6B70;
}

html, body, .stApp, [data-testid="stAppViewContainer"] {
  font-family: 'DM Sans', sans-serif;
  color: var(--ink);
}
.stApp { background: var(--paper); }
#MainMenu, footer { visibility: hidden; }
header[data-testid="stHeader"] { background: transparent; }
.block-container { max-width: 820px; padding-top: 2.2rem; padding-bottom: 4rem; }

h1, h2, h3, h4, h5 {
  font-family: 'Bricolage Grotesque', sans-serif;
  letter-spacing: -0.01em;
  color: var(--ink);
}

/* En-tête de marque */
.brand { display: flex; align-items: center; gap: 16px; margin: 0 0 1.6rem; }
.brand svg { width: 58px; height: 58px; flex: none; }
.brand h1 { font-size: 2rem; line-height: 1.1; margin: 0; padding: 0; }
.brand p { margin: .25rem 0 0; color: var(--muted); font-size: 1rem; }

/* Formulaire */
[data-testid="stForm"] {
  background: #FFFFFF;
  border: 1px solid var(--line);
  border-radius: 14px;
  padding: 1.5rem 1.5rem 1.1rem;
}
[data-testid="stForm"] h5 { margin: .4rem 0 .2rem; font-size: 1.05rem; }
.stTextInput input, .stTextArea textarea { border-radius: 10px; }
[data-baseweb="select"] > div { border-radius: 10px; }

/* Boutons */
.stFormSubmitButton button, .stButton button, .stDownloadButton button {
  border-radius: 10px;
  font-weight: 700;
  padding: .65rem 1rem;
}
.stFormSubmitButton button { width: 100%; }
.stButton, .stDownloadButton, .stLinkButton, [data-testid="stPopover"] { width: 100%; }
.stButton button, .stDownloadButton button, .stLinkButton a, [data-testid="stPopover"] button { width: 100%; }

/* Barre latérale */
[data-testid="stSidebar"] { background: #FFFFFF; border-right: 1px solid var(--line); }
.tip { color: var(--muted); font-size: .92rem; line-height: 1.5; }

/* Résultat */
.result-meta { color: var(--muted); font-size: .92rem; margin: -.2rem 0 .6rem; }

@media (prefers-reduced-motion: reduce) {
  * { animation: none !important; transition: none !important; }
}
</style>
"""

st.markdown(CSS, unsafe_allow_html=True)


# ─────────────────────────────────────────────────────────────
# Statistiques réelles (persistées dans un fichier JSON)
# ─────────────────────────────────────────────────────────────
STATS_FILE = Path("stats.json")
_STATS_LOCK = threading.Lock()
DEFAULT_STATS = {"likes": 0, "generations": 0}


def load_stats() -> dict:
    try:
        data = json.loads(STATS_FILE.read_text(encoding="utf-8"))
        return {k: int(data.get(k, v)) for k, v in DEFAULT_STATS.items()}
    except Exception:  # noqa: BLE001
        return dict(DEFAULT_STATS)


def bump_stat(key: str) -> dict:
    """Incrémente un compteur de façon sûre (verrou + écriture tolérante aux erreurs)."""
    with _STATS_LOCK:
        stats = load_stats()
        stats[key] = stats.get(key, 0) + 1
        try:
            STATS_FILE.write_text(json.dumps(stats), encoding="utf-8")
        except OSError:
            pass  # système de fichiers en lecture seule : on ignore
        return stats


if "user_has_liked" not in st.session_state:
    st.session_state.user_has_liked = False


# ─────────────────────────────────────────────────────────────
# Utilitaires
# ─────────────────────────────────────────────────────────────
def get_secret(name: str, default: str = "") -> str:
    """Lit une valeur dans les secrets Streamlit sans planter si absents."""
    try:
        return str(st.secrets.get(name, default) or default)
    except Exception:  # noqa: BLE001
        return default


def slugify(value: str) -> str:
    return re.sub(r"[^\w-]+", "_", value.strip(), flags=re.UNICODE).strip("_") or "document"


def clean_output(text: str) -> str:
    """Retire le Markdown résiduel pour obtenir un texte brut réutilisable."""
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"__(.*?)__", r"\1", text)
    text = re.sub(r"^#{1,6}\s*", "", text, flags=re.M)
    text = re.sub(r"^\s*[\*\-•]\s+", "• ", text, flags=re.M)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def friendly_error(err: Exception) -> str:
    msg = str(err).lower()
    # Expressions précises : « rate » seul apparaît dans « generate », d'où de faux positifs.
    if re.search(r"api key|api_key|permission|unauthenticated|\b40[13]\b", msg):
        return "La clé API est refusée. Vérifiez qu'elle est correcte et active."
    if re.search(r"\b429\b|quota|rate limit|resource_exhausted", msg):
        return "La limite d'utilisation de l'API est atteinte. Réessayez dans une minute."
    if re.search(r"\b503\b|overloaded|unavailable", msg):
        return "Le service de rédaction est saturé. Réessayez dans quelques instants."
    return f"La génération a échoué. Détail technique : {err}"


STOPWORDS = set(
    """
    dans pour avec vous nous votre notre vos nos être sera seront cette cettes entre leurs leur
    aussi ainsi alors plus moins comme chez sous sont elle elles ils dont mais donc car
    poste offre mission missions profil candidat entreprise société travail équipe
    recherchons rejoindre rejoignez avoir faire tous toutes tout toute très bien
    """.split()
)


def extract_keywords(text: str, top: int = 20) -> list[str]:
    words = re.findall(r"[a-zà-ÿœ][a-zà-ÿœ'-]{4,}", text.lower())
    counts = Counter(w for w in words if w not in STOPWORDS)
    return [w for w, _ in counts.most_common(top)]


def keyword_match(offer: str, document: str):
    keywords = extract_keywords(offer)
    if not keywords:
        return None
    doc = document.lower()
    missing = [k for k in keywords if k not in doc]
    return len(keywords) - len(missing), len(keywords), missing


# ─────────────────────────────────────────────────────────────
# Génération IA
# ─────────────────────────────────────────────────────────────
def build_prompts(p: dict) -> tuple[str, str]:
    is_cv = p["doc_type"].startswith("CV")
    kind = "cv" if is_cv else "lettre"
    length = LENGTHS[kind][p["length"]]

    if is_cv:
        structure = (
            "STRUCTURE DU CV (chaque titre de section en MAJUSCULES, seul sur sa ligne) :\n"
            "- Ligne 1 : nom et prénom. Ligne 2 : intitulé du poste visé.\n"
            "- PROFIL : 3 lignes maximum.\n"
            "- COMPÉTENCES CLÉS : liste courte, orientée vers l'offre.\n"
            "- EXPÉRIENCES PROFESSIONNELLES : pour chaque expérience, une ligne "
            "« Poste, Structure (période) » puis des puces commençant par un verbe d'action "
            "et, si possible, un résultat.\n"
            "- FORMATION\n"
            "- LANGUES ET ATOUTS (si pertinent)\n"
            f"Le CV {length}."
        )
    else:
        structure = (
            "STRUCTURE DE LA LETTRE :\n"
            "- Coordonnées du candidat (entre crochets si inconnues), lieu et date, "
            "destinataire.\n"
            "- Objet.\n"
            "- Formule d'appel.\n"
            "- 3 ou 4 paragraphes : accroche personnalisée, valeur ajoutée du candidat "
            "avec des exemples concrets, motivation pour l'entreprise, ouverture vers un entretien.\n"
            "- Formule de politesse et signature.\n"
            f"La lettre fait {length}."
        )

    system = (
        "Tu es un expert en recrutement et en rédaction de CV et de lettres de motivation, "
        "capable d'écrire des documents optimisés pour les logiciels de tri (ATS) tout en "
        "restant agréables à lire pour un recruteur.\n"
        "RÈGLES ABSOLUES :\n"
        "1. N'invente JAMAIS de fait : employeur, diplôme, date, chiffre ou outil absent des "
        "informations fournies. Si une information utile manque, insère un champ entre "
        "crochets, par exemple [ville] ou [chiffre à préciser].\n"
        "2. Écris en texte brut : pas de Markdown, pas de titres avec #, pas de gras. "
        "Utilise « • » pour les puces.\n"
        "3. Style clair, concret, sans formules creuses ni superlatifs inutiles.\n"
        f"4. Rédige en {p['language']}.\n"
        "5. Réponds uniquement avec le document final, sans commentaire avant ni après."
    )

    offer = p["offer"] or "Aucune offre fournie. Base-toi sur les standards du poste visé."
    prompt = f"""Rédige : {p['doc_type']}

INFORMATIONS SUR LE CANDIDAT
- Nom et prénom : {p['name']}
- Poste visé : {p['job']}
- Entreprise ciblée : {p['company'] or 'Non précisée'}
- Parcours et compétences : {p['background']}
- Ton souhaité : {p['tone']}
- Consignes particulières : {p['notes'] or 'Aucune'}

OFFRE D'EMPLOI DE RÉFÉRENCE
{offer}

{structure}

Si une offre est fournie, reprends naturellement ses mots-clés et compétences attendues, \
uniquement lorsqu'ils correspondent au parcours du candidat."""
    return system, prompt


def stream_document(api_key: str, system: str, prompt: str, temperature: float):
    """Génère le texte en flux continu, avec un modèle de secours si le premier échoue."""
    client = genai.Client(api_key=api_key)
    last_error = None
    for model in MODELS:
        started = False
        try:
            stream = client.models.generate_content_stream(
                model=model,
                contents=prompt,
                config=types.GenerateContentConfig(
                    system_instruction=system,
                    temperature=temperature,
                ),
            )
            for chunk in stream:
                if chunk.text:
                    started = True
                    yield chunk.text
            if started:
                return
            last_error = RuntimeError("Réponse vide du modèle.")
        except Exception as err:  # noqa: BLE001
            if started:
                raise  # flux interrompu en cours de route : inutile de relancer
            last_error = err
    raise last_error


# ─────────────────────────────────────────────────────────────
# Export PDF
# ─────────────────────────────────────────────────────────────
INK = (18, 39, 43)
BODY = (38, 52, 56)
TEAL = (14, 107, 107)
AMBER = (242, 169, 59)
LINE = (220, 227, 229)

PDF_REPLACEMENTS = {
    "\u2019": "'", "\u2018": "'", "\u201c": '"', "\u201d": '"', "\u2013": "-", "\u2014": "-",
    "\u2026": "...", "\u2022": "-", "\u00a0": " ", "\u202f": " ", "\u0153": "oe", "\u0152": "Oe",
    "\u20ac": "EUR", "\u2192": "->",
}


def pdf_safe(text: str) -> str:
    for src, dst in PDF_REPLACEMENTS.items():
        text = text.replace(src, dst)
    text = text.encode("latin-1", "replace").decode("latin-1")
    # Coupe les « mots » très longs (URL…) qui feraient planter multi_cell.
    return re.sub(r"(\S{60})(?=\S)", r"\1 ", text)


def is_heading(line: str) -> bool:
    core = line.rstrip(":").strip()
    letters = [c for c in core if c.isalpha()]
    return bool(letters) and core == core.upper() and len(core) <= 48 and not core.startswith("-")


@st.cache_data(show_spinner=False)
def build_pdf(text: str, is_cv: bool) -> bytes:
    pdf = FPDF(format="A4")
    pdf.set_margins(20, 18, 20)
    pdf.set_auto_page_break(True, margin=18)
    pdf.add_page()
    right = pdf.w - pdf.r_margin
    header_step = 0

    for raw in text.split("\n"):
        line = pdf_safe(raw).strip()
        if not line:
            pdf.ln(3)
            continue

        if is_cv and header_step == 0:
            pdf.set_font("Helvetica", "B", 22)
            pdf.set_text_color(*INK)
            pdf.multi_cell(0, 10, line, new_x="LMARGIN", new_y="NEXT")
            header_step = 1
            continue

        if is_cv and header_step == 1:
            header_step = 2
            if not is_heading(line):
                pdf.set_font("Helvetica", "", 12)
                pdf.set_text_color(*TEAL)
                pdf.multi_cell(0, 7, line, new_x="LMARGIN", new_y="NEXT")
                y = pdf.get_y() + 2
                pdf.set_draw_color(*AMBER)
                pdf.set_line_width(0.8)
                pdf.line(pdf.l_margin, y, pdf.l_margin + 40, y)
                pdf.ln(5)
                continue

        if is_cv and is_heading(line):
            pdf.ln(3)
            pdf.set_font("Helvetica", "B", 11)
            pdf.set_text_color(*TEAL)
            pdf.multi_cell(0, 6, line.rstrip(":"), new_x="LMARGIN", new_y="NEXT")
            y = pdf.get_y() + 0.5
            pdf.set_draw_color(*LINE)
            pdf.set_line_width(0.3)
            pdf.line(pdf.l_margin, y, right, y)
            pdf.ln(2.5)
        elif line.startswith("- "):
            pdf.set_font("Helvetica", "", 10.5)
            pdf.set_text_color(*BODY)
            pdf.set_x(pdf.l_margin + 2)
            pdf.cell(5, 5.6, "-")
            pdf.multi_cell(0, 5.6, line[2:].strip(), new_x="LMARGIN", new_y="NEXT")
        else:
            pdf.set_font("Helvetica", "", 10.5)
            pdf.set_text_color(*BODY)
            pdf.multi_cell(0, 5.6, line, new_x="LMARGIN", new_y="NEXT")

    return bytes(pdf.output())


# ─────────────────────────────────────────────────────────────
# Export Word (.docx)
# ─────────────────────────────────────────────────────────────
BULLET_RE = re.compile(r"^[•\-\*]\s+")


def _rgb(color: tuple):
    return RGBColor(*color)


def _bottom_border(paragraph, color: str = "DCE3E5", size: int = 6) -> None:
    """Ajoute un filet sous un paragraphe (titres de section)."""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


@st.cache_data(show_spinner=False)
def build_docx(text: str, is_cv: bool) -> bytes:
    """Génère un document Word A4 mis en forme, modifiable dans Word / LibreOffice / Google Docs."""
    doc = Document()

    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = section.bottom_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = _rgb(BODY)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.1

    header_step = 0

    for raw in text.split("\n"):
        line = raw.strip()

        if not line:
            if not is_cv:  # lettre : les lignes vides séparent les paragraphes
                doc.add_paragraph().paragraph_format.space_after = Pt(0)
            continue

        if is_cv and header_step == 0:
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.bold = True
            run.font.size = Pt(22)
            run.font.color.rgb = _rgb(INK)
            p.paragraph_format.space_after = Pt(0)
            header_step = 1
            continue

        if is_cv and header_step == 1:
            header_step = 2
            if not is_heading(line):
                p = doc.add_paragraph()
                run = p.add_run(line)
                run.font.size = Pt(12.5)
                run.font.color.rgb = _rgb(TEAL)
                p.paragraph_format.space_after = Pt(6)
                _bottom_border(p, color="F2A93B", size=12)
                continue

        if is_cv and is_heading(line):
            p = doc.add_paragraph()
            run = p.add_run(line.rstrip(":"))
            run.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = _rgb(TEAL)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            _bottom_border(p)
        elif BULLET_RE.match(line):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(BULLET_RE.sub("", line))
            p.paragraph_format.space_after = Pt(1.5)
        else:
            p = doc.add_paragraph(line)
            if not is_cv:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(4)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ─────────────────────────────────────────────────────────────
# Interface
# ─────────────────────────────────────────────────────────────
st.markdown(
    f'<div class="brand">{LOGO_SVG}<div><h1>Pavel IA CV</h1>'
    "<p>Un CV ou une lettre de motivation adaptés à l'offre, prêts à télécharger.</p></div></div>",
    unsafe_allow_html=True,
)

with st.sidebar:
    st.markdown("### Réglages")
    api_key = get_secret("GEMINI_API_KEY")
    if not api_key:
        api_key = st.text_input(
            "Clé API Gemini",
            type="password",
            help="Clé gratuite sur https://aistudio.google.com/",
        ).strip()
    creativity = st.slider(
        "Créativité",
        0.0, 1.0, 0.6, 0.1,
        help="Bas : texte sobre et fidèle à vos données. Haut : formulations plus libres.",
    )
    st.markdown(
        '<p class="tip">Plus vous donnez de faits précis (chiffres, outils, résultats), '
        "plus le document est convaincant. Collez l'offre complète pour que les mots-clés "
        "soient repris.</p>",
        unsafe_allow_html=True,
    )

    st.divider()

    # --- SECTION STATISTIQUES (réelles) ---
    stats = load_stats()
    st.markdown("### Statistiques")
    col_stat1, col_stat2 = st.columns(2)
    with col_stat1:
        st.metric(label="Docs créés", value=stats["generations"])
    with col_stat2:
        st.metric(label="Recommandations", value=stats["likes"])

   
